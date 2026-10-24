#!/usr/bin/env python3
"""
Render docs/outreach/SUMMARY.md to docs/outreach/SUMMARY.docx (one page).

SUMMARY.md is the source of truth; this script parses its lightweight Markdown
(H1 title, **bold lead-in:** paragraphs, *italic* spans, and "- " bullets) so the
two stay in sync. Every sentence is taken from SUMMARY.md.

Requires python-docx (pip install python-docx).
Run:  python3 build_summary_docx.py
"""
import os
import re
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "SUMMARY.md")
OUT = os.path.join(HERE, "SUMMARY.docx")

TOKEN = re.compile(r"(\*\*.+?\*\*|\*.+?\*)")


def add_runs(paragraph, text):
    """Render inline **bold** and *italic* spans into runs (no nesting)."""
    for piece in TOKEN.split(text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**"):
            r = paragraph.add_run(piece[2:-2]); r.bold = True
        elif piece.startswith("*") and piece.endswith("*"):
            r = paragraph.add_run(piece[1:-1]); r.italic = True
        else:
            paragraph.add_run(piece)


def main():
    lines = open(SRC, encoding="utf-8").read().splitlines()

    doc = Document()
    # tight one-page geometry
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Inches(0.7)
        s.left_margin = s.right_margin = Inches(0.85)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.04

    i = 0
    n = len(lines)
    while i < n:
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
            continue
        if line.startswith("# "):
            h = doc.add_paragraph()
            r = h.add_run(line[2:].strip())
            r.bold = True
            r.font.size = Pt(17)
            r.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
            h.paragraph_format.space_after = Pt(8)
        elif line.startswith("- "):
            # consume a contiguous bullet block
            while i < n and lines[i].rstrip().startswith("- "):
                btxt = lines[i].rstrip()[2:].strip()
                # join wrapped continuation lines (indented, not blank, not a new bullet)
                j = i + 1
                while j < n and lines[j].startswith("  ") and lines[j].strip() and not lines[j].lstrip().startswith("- "):
                    btxt += " " + lines[j].strip()
                    j += 1
                p = doc.add_paragraph(style="List Bullet")
                p.paragraph_format.space_after = Pt(3)
                add_runs(p, btxt)
                i = j
            continue
        else:
            # paragraph: join wrapped lines until blank
            para = line
            j = i + 1
            while j < n and lines[j].strip() and not lines[j].startswith(("# ", "- ")):
                para += " " + lines[j].strip()
                j += 1
            p = doc.add_paragraph()
            add_runs(p, para)
            i = j
            continue
        i += 1

    doc.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    main()
